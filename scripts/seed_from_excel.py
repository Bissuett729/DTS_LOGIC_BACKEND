"""
Seeder: Down Time February 2026
================================
Lee el archivo Excel "Down Time February 2026 (2).xlsx" y carga todos
los registros a la base de datos a través del API DTS.

Cada fila del Excel = un registro de 1 hora en la colección "Down Time".

Columnas del Excel:
  A  DATE            → startTime (fecha + hora de inicio)
  B  WK              → week
  C  SHIFT           → shift (A/B/C/D)
  D  LINE            → line
  E  START TIME      → startTime (hora)
  F  FINISH TIME     → endTime  (hora + 1)
  G  ID SUPERVISOR   → supervisor
  H  STD             → standardOutput
  I  REAL            → currentOutput
  J  EFFIC           → efficiency
  K  DOWN TIME NOT REPORTED → downTimeUnreported
  L  DOWN TIME GENERATED    → downTimeGenerated
  M-U  Departamentos (TEST, FACILITIES/ICC, PMC, AUTO, MFG, PROCESS, QA, IT/FISS/SIE, TRAINING/RH)
  V  TOTAL DOWNTIME REPORTED → downTimeReported
  W  PROBLEM 1  /  X  ACTION 1
  Y  PROBLEM 2  /  Z  ACTION 2
  AA PROBLEM 3  /  AB ACTION 3
  AC PROBLEM 4  /  AD ACTION 4

Uso:
  python seed_from_excel.py [--dry-run] [--sheet "FA DOCKING 1"] [--delay 0.05]
"""

import argparse
import datetime
import os
import sys
import time

import openpyxl
import requests

# ─────────────────────────────────────────────
# CONFIGURACIÓN
# ─────────────────────────────────────────────
EXCEL_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "../../DTS_FRONTEND/src/app/features/downtime-register/Down Time February 2026 (2).xlsx",
)
API_URL = "http://10.19.16.37:20026/v1/down-time"
REGISTERED_BY = "seeder"
# "DATA NOVEMBER" y "DATA" son sheets consolidados que duplican los
# datos de los 33 sheets individuales (33 × 744 = 24,552 filas cada uno).
# "Causas" es hoja de referencia sin registros de operación.
SKIP_SHEETS = {"Sheet1", "Causas", "DATA NOVEMBER", "DATA"}

# Índice 0-based de columnas en la fila de valores
COL_DATE       = 0
COL_WEEK       = 1
COL_SHIFT      = 2
COL_LINE       = 3
COL_START_HOUR = 4
COL_END_HOUR   = 5
COL_SUPERVISOR = 6
COL_STD        = 7
COL_REAL       = 8
COL_EFFIC      = 9
COL_DT_NOT_REP = 10
COL_DT_GEN     = 11
# Columnas M-U (índice 12–20) = departamentos
COL_DEPTS_START = 12
COL_DEPTS_END   = 21   # exclusivo (9 columnas)
COL_TOTAL_REP   = 21
# Columnas W, Y, AA, AC = problemas (índice 22, 24, 26, 28)
COL_PROBLEMS = [22, 24, 26, 28]


# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────
def get_stage(sheet_name: str) -> str:
    """Obtiene el stage del nombre del sheet (primera palabra)."""
    return sheet_name.split(" ")[0]


def to_float(val) -> float | None:
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def to_int(val) -> int | None:
    f = to_float(val)
    return int(f) if f is not None else None


def build_datetime(base_date: datetime.datetime, hour: int) -> str:
    """Combina la fecha base con la hora y retorna ISO-8601 string."""
    # Si hour >= 24 pasa al día siguiente
    delta_days = hour // 24
    hour_norm  = hour % 24
    dt = base_date.replace(hour=hour_norm, minute=0, second=0, microsecond=0)
    dt += datetime.timedelta(days=delta_days)
    return dt.strftime("%Y-%m-%dT%H:%M:%S.000Z")


def build_classification(row, dept_headers: list[str]) -> list[dict]:
    """Construye la lista de clasificaciones desde las columnas M-U."""
    classification = []
    problems = []
    for prob_col in COL_PROBLEMS:
        val = row[prob_col] if prob_col < len(row) else None
        if val and str(val).strip():
            problems.append(str(val).strip())

    prob_idx = 0
    for i in range(9):  # 9 departamentos
        col_idx = COL_DEPTS_START + i
        val = row[col_idx] if col_idx < len(row) else None
        minutes = to_int(val)
        if minutes and minutes > 0:
            dept = dept_headers[i].strip() if i < len(dept_headers) and dept_headers[i] else f"Dept{i+1}"
            # Normalizar nombre de departamento
            dept = dept.replace("/ ", "/").replace("IT/FISS", "IT/FISS").strip()
            reason = problems[prob_idx] if prob_idx < len(problems) else ""
            prob_idx += 1
            classification.append(
                {"downTimeGenerated": minutes, "department": dept, "reason": reason}
            )
    return classification


def build_payload(row, dept_headers: list[str], stage: str) -> dict | None:
    """Construye el payload para el API. Retorna None si la fila no es válida."""
    base_date = row[COL_DATE]
    if not isinstance(base_date, datetime.datetime):
        return None

    start_hour = to_int(row[COL_START_HOUR])
    end_hour   = to_int(row[COL_END_HOUR])
    if start_hour is None or end_hour is None:
        return None

    week       = to_int(row[COL_WEEK])
    shift      = str(row[COL_SHIFT]).strip() if row[COL_SHIFT] else None
    line       = str(row[COL_LINE]).strip() if row[COL_LINE] else None
    supervisor = to_int(row[COL_SUPERVISOR])
    std        = to_int(row[COL_STD])
    real       = to_int(row[COL_REAL])
    effic      = to_float(row[COL_EFFIC])
    dt_gen     = to_int(row[COL_DT_GEN])
    dt_not_rep = to_int(row[COL_DT_NOT_REP])
    dt_rep     = to_int(row[COL_TOTAL_REP])

    classification = build_classification(row, dept_headers)

    payload: dict = {
        "startTime":        build_datetime(base_date, start_hour),
        "endTime":          build_datetime(base_date, end_hour),
        "registeredBy":     REGISTERED_BY,
    }

    if week        is not None: payload["week"]            = week
    if shift:                   payload["shift"]           = shift
    if line:                    payload["line"]            = line
    if stage:                   payload["stage"]           = stage
    if supervisor  is not None: payload["supervisor"]      = str(supervisor)
    if std         is not None: payload["standardOutput"]  = std
    if real        is not None: payload["currentOutput"]   = real
    if effic       is not None: payload["efficiency"]      = round(effic, 4)
    if dt_gen      is not None: payload["downTimeGenerated"]   = dt_gen
    if dt_not_rep  is not None: payload["downTimeUnreported"]  = dt_not_rep
    if dt_rep      is not None: payload["downTimeReported"]    = dt_rep
    if classification:          payload["classification"]  = classification

    return payload


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Seeder: Down Time Excel → API")
    parser.add_argument("--dry-run", action="store_true", help="Solo parsea, no envía al API")
    parser.add_argument("--sheet",   type=str, default=None, help="Procesar solo este sheet")
    parser.add_argument("--delay",   type=float, default=0.0, help="Segundos entre peticiones (default 0)")
    parser.add_argument("--limit",   type=int, default=None,  help="Máximo de registros a enviar (para pruebas)")
    args = parser.parse_args()

    excel_path = os.path.normpath(EXCEL_PATH)
    if not os.path.exists(excel_path):
        print(f"❌  No se encontró el archivo: {excel_path}")
        sys.exit(1)

    print(f"📂  Leyendo: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    sheets = wb.sheetnames
    if args.sheet:
        if args.sheet not in sheets:
            print(f"❌  Sheet '{args.sheet}' no encontrado. Disponibles: {sheets}")
            sys.exit(1)
        sheets = [args.sheet]

    total_ok    = 0
    total_skip  = 0
    total_error = 0
    total_limit_reached = False

    for sheet_name in sheets:
        if sheet_name in SKIP_SHEETS:
            continue

        ws = wb[sheet_name]
        stage = get_stage(sheet_name)

        # Leer headers de la fila 1 para obtener nombres de departamentos (M-U)
        header_row = [cell.value for cell in ws[1]]
        dept_headers = header_row[COL_DEPTS_START:COL_DEPTS_END]

        print(f"\n📋  Sheet: {sheet_name}  (stage={stage})")
        print(f"    Departamentos: {dept_headers}")

        sheet_ok = 0
        sheet_skip = 0

        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if args.limit and (total_ok + total_error) >= args.limit:
                total_limit_reached = True
                break

            payload = build_payload(row, dept_headers, stage)

            if payload is None:
                sheet_skip += 1
                total_skip += 1
                continue

            if args.dry_run:
                print(f"  [DRY] row {row_num}: {payload}")
                sheet_ok += 1
                total_ok += 1
                continue

            try:
                resp = requests.post(API_URL, json=payload, timeout=15)
                if resp.status_code in (200, 201):
                    sheet_ok += 1
                    total_ok += 1
                    if sheet_ok % 50 == 0:
                        print(f"    ✅  {sheet_ok} registros enviados en este sheet...")
                else:
                    total_error += 1
                    print(f"  ⚠️   row {row_num} HTTP {resp.status_code}: {resp.text[:300]}")
            except requests.exceptions.ConnectionError:
                total_error += 1
                print(f"  ❌  row {row_num}: No se pudo conectar con {API_URL}")
                print("       Verifica que el backend esté corriendo.")
                sys.exit(1)
            except requests.exceptions.Timeout:
                total_error += 1
                print(f"  ⏱️   row {row_num}: Timeout")

            if args.delay > 0:
                time.sleep(args.delay)

        print(f"    → Enviados: {sheet_ok}  |  Omitidos: {sheet_skip}")

        if total_limit_reached:
            print(f"\n⚠️  Límite de {args.limit} registros alcanzado.")
            break

    print("\n" + "═" * 50)
    print(f"✅  Registros insertados : {total_ok}")
    print(f"⏭️   Filas omitidas       : {total_skip}")
    print(f"❌  Errores              : {total_error}")
    print("═" * 50)


if __name__ == "__main__":
    main()
